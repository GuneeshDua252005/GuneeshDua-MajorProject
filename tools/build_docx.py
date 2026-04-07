#!/usr/bin/env python3
"""
Convert markdown research paper to DOCX deliverable.
Lightweight converter for project packaging purposes.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
INPUT_MD = ROOT / "docs" / "IEEE_RESEARCH_PAPER_CEI.md"
OUTPUT_DOCX = ROOT / "deliverables" / "CEI_Research_Paper_IEEE.docx"


def is_horizontal_rule(line: str) -> bool:
    stripped = line.strip()
    return stripped in {"---", "***", "___"}


def add_markdown_line(doc: Document, line: str) -> None:
    stripped = line.rstrip()
    if not stripped:
        doc.add_paragraph("")
        return
    if is_horizontal_rule(stripped):
        doc.add_paragraph("")
        return
    if stripped.startswith("### "):
        doc.add_heading(stripped[4:].strip(), level=3)
        return
    if stripped.startswith("## "):
        doc.add_heading(stripped[3:].strip(), level=2)
        return
    if stripped.startswith("# "):
        doc.add_heading(stripped[2:].strip(), level=1)
        return
    if stripped.startswith("|") and stripped.endswith("|"):
        # Keep markdown table rows as monospace lines for readability in DOCX.
        p = doc.add_paragraph(stripped)
        if p.runs:
            p.runs[0].font.name = "Consolas"
            p.runs[0].font.size = Pt(9)
        return
    if stripped.startswith("```"):
        return
    if stripped.startswith("- "):
        doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
        return
    if stripped[:2].isdigit() and stripped[2:4] in {". ", ") "}:
        doc.add_paragraph(stripped[4:].strip(), style="List Number")
        return
    if stripped.startswith("**") and stripped.endswith("**") and len(stripped) > 4:
        p = doc.add_paragraph(stripped.strip("*"))
        if p.runs:
            p.runs[0].bold = True
        return
    doc.add_paragraph(stripped)


def build_docx() -> None:
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    if not INPUT_MD.exists():
        raise FileNotFoundError(f"Input markdown not found: {INPUT_MD}")

    text = INPUT_MD.read_text(encoding="utf-8").splitlines()
    doc = Document()
    doc.core_properties.title = "Cognitive Emotion Intelligence and Adaptive Lifestyle System"
    doc.core_properties.subject = "IEEE Style Major Project Research Paper"
    doc.core_properties.author = "Guneesh Dua"

    for line in text:
        add_markdown_line(doc, line)

    doc.save(OUTPUT_DOCX)
    print(f"DOCX created at: {OUTPUT_DOCX}")


if __name__ == "__main__":
    build_docx()
