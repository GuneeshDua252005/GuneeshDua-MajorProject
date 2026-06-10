#!/usr/bin/env python3
"""
Generate a .docx version of the research paper from docs/research_paper.md.

This script uses python-docx when available. It applies a simple academic
heading structure and preserves code blocks in monospace paragraphs.
"""

from __future__ import annotations

from pathlib import Path


def parse_markdown_sections(markdown_text: str):
    in_code_block = False
    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip("\n")
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code_block = not in_code_block
            continue
        yield line, in_code_block


def main() -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
        from docx.shared import Pt
    except ImportError as exc:
        raise SystemExit(
            "python-docx is required. Install dependencies with: pip install -r requirements.txt"
        ) from exc

    root = Path(__file__).resolve().parents[1]
    source_path = root / "docs" / "research_paper.md"
    output_dir = root / "deliverables"
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "Cognitive_Emotion_Intelligence_Adaptive_Lifestyle_System_IEEE_Paper.docx"

    source_text = source_path.read_text(encoding="utf-8")

    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    title = document.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = title.add_run("Cognitive Emotion Intelligence and Adaptive Lifestyle System")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = subtitle.add_run("IEEE-Style Research Paper Draft")
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)

    document.add_paragraph("")

    for line, in_code_block in parse_markdown_sections(source_text):
        stripped = line.strip()
        if not stripped:
            document.add_paragraph("")
            continue

        if stripped.startswith("# "):
            p = document.add_paragraph()
            r = p.add_run(stripped[2:].strip())
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(13)
            continue

        if stripped.startswith("## "):
            p = document.add_paragraph()
            r = p.add_run(stripped[3:].strip())
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            continue

        if stripped.startswith("### "):
            p = document.add_paragraph()
            r = p.add_run(stripped[4:].strip())
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            continue

        if in_code_block:
            p = document.add_paragraph()
            r = p.add_run(line)
            r.font.name = "Courier New"
            r.font.size = Pt(9.5)
            continue

        if stripped.startswith("- "):
            p = document.add_paragraph(style="List Bullet")
            r = p.add_run(stripped[2:].strip())
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            continue

        document.add_paragraph(line)

    document.save(output_path)
    print(output_path)


if __name__ == "__main__":
    main()
