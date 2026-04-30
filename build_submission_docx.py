from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
INPUT_MD = ROOT / "GGSIPU_Major_Project_Report_CEIALOS.md"
OUTPUT_DOCX = ROOT / "GGSIPU_Major_Project_Report_CEIALOS.docx"
DEPARTMENT_FOOTER = "NAME OF THE DEPARTMENT"


def set_page_number_format(section, fmt: str, start: int = 1) -> None:
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:fmt"), fmt)
    pg_num_type.set(qn("w:start"), str(start))


def add_page_field(paragraph) -> None:
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def configure_footer(section, show_footer: bool = True) -> None:
    section.footer.is_linked_to_previous = False
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.clear()
    if not show_footer:
        return

    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.0))
    left_run = p.add_run(DEPARTMENT_FOOTER)
    left_run.font.name = "Times New Roman"
    left_run.font.size = Pt(10)
    p.add_run("\t")
    add_page_field(p)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.different_first_page_header_footer = True

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, bold, italic in [
        ("Heading 1", 16, True, False),
        ("Heading 2", 14, True, False),
        ("Heading 3", 12, False, True),
    ]:
        h = doc.styles[name]
        h.font.name = "Times New Roman"
        h.font.size = Pt(size)
        h.font.bold = bold
        h.font.italic = italic
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h.paragraph_format.line_spacing = 1.5

    if "CodeBlock" not in [s.name for s in doc.styles]:
        code_style = doc.styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = "Courier New"
        code_style.font.size = Pt(10)

    set_page_number_format(section, fmt="lowerRoman", start=1)
    configure_footer(section, show_footer=True)
    return doc


def parse_table_row(line: str) -> list[str]:
    parts = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return parts


def is_separator_row(cells: list[str]) -> bool:
    return all(re.fullmatch(r":?-{2,}:?", c or "") is not None for c in cells)


def add_markdown_table(doc: Document, rows: list[str]) -> None:
    parsed = [parse_table_row(r) for r in rows if r.strip()]
    if len(parsed) < 2:
        return
    headers = parsed[0]
    body = []
    for row in parsed[1:]:
        if len(row) == len(headers) and not is_separator_row(row):
            body.append(row)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = text
    for row in body:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text


def convert_markdown_to_docx(markdown_text: str, doc: Document) -> None:
    lines = markdown_text.splitlines()
    in_code = False
    table_buffer: list[str] = []
    first_h1_seen = False
    chapter_one_started = False

    i = 0
    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("|"):
            table_buffer.append(line)
            i += 1
            continue
        if table_buffer:
            add_markdown_table(doc, table_buffer)
            table_buffer.clear()

        if line.strip().startswith("```"):
            in_code = not in_code
            i += 1
            continue

        if in_code:
            p = doc.add_paragraph(style="CodeBlock")
            p.add_run(line)
            i += 1
            continue

        h_match = re.match(r"^(#{1,3})\s+(.*)$", line.strip())
        if h_match:
            level = len(h_match.group(1))
            title = h_match.group(2).strip()

            if level == 1:
                if first_h1_seen:
                    if title.upper().startswith("CHAPTER 1:") and not chapter_one_started:
                        new_sec = doc.add_section(WD_SECTION.NEW_PAGE)
                        new_sec.left_margin = Inches(1.5)
                        new_sec.right_margin = Inches(1.0)
                        new_sec.top_margin = Inches(1.0)
                        new_sec.bottom_margin = Inches(1.0)
                        new_sec.different_first_page_header_footer = False
                        set_page_number_format(new_sec, fmt="decimal", start=1)
                        configure_footer(new_sec, show_footer=True)
                        chapter_one_started = True
                    else:
                        doc.add_page_break()
                else:
                    first_h1_seen = True

            doc.add_paragraph(title, style=f"Heading {level}")
            i += 1
            continue

        if not line.strip():
            doc.add_paragraph("")
            i += 1
            continue

        if re.match(r"^\s*[-*]\s+", line):
            text = re.sub(r"^\s*[-*]\s+", "", line).strip()
            doc.add_paragraph(text, style="List Bullet")
            i += 1
            continue

        if re.match(r"^\s*\d+\.\s+", line):
            text = re.sub(r"^\s*\d+\.\s+", "", line).strip()
            doc.add_paragraph(text, style="List Number")
            i += 1
            continue

        p = doc.add_paragraph(line.strip(), style="Normal")
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        i += 1

    if table_buffer:
        add_markdown_table(doc, table_buffer)


def main() -> None:
    if not INPUT_MD.exists():
        raise FileNotFoundError(f"Missing input markdown: {INPUT_MD}")
    markdown_text = INPUT_MD.read_text(encoding="utf-8")
    doc = setup_document()
    convert_markdown_to_docx(markdown_text, doc)
    doc.save(OUTPUT_DOCX)
    print(f"DOCX generated: {OUTPUT_DOCX}")


if __name__ == "__main__":
    main()
