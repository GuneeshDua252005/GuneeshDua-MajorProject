from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "ieee_research_paper.md"
OUTPUT = ROOT / "Cognitive_Emotion_Intelligence_Adaptive_Lifestyle_System_IEEE_Paper.docx"


def add_paragraph_from_markdown(document: Document, line: str) -> None:
    stripped = line.strip()
    if not stripped:
        document.add_paragraph("")
        return

    if stripped.startswith("# "):
        document.add_heading(stripped[2:].strip(), level=1)
        return
    if stripped.startswith("## "):
        document.add_heading(stripped[3:].strip(), level=2)
        return
    if stripped.startswith("### "):
        document.add_heading(stripped[4:].strip(), level=3)
        return
    if stripped.startswith("#### "):
        document.add_heading(stripped[5:].strip(), level=4)
        return
    if stripped.startswith("- "):
        document.add_paragraph(stripped[2:].strip(), style="List Bullet")
        return
    if stripped[:2].isdigit() and stripped[1:3] == ". ":
        document.add_paragraph(stripped[3:].strip(), style="List Number")
        return
    if stripped.startswith("```"):
        return
    if stripped == "---":
        document.add_paragraph("")
        return
    document.add_paragraph(stripped)


def generate() -> Path:
    text = SOURCE.read_text(encoding="utf-8")
    document = Document()

    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    for section in document.sections:
        section.top_margin = Pt(54)
        section.bottom_margin = Pt(54)
        section.left_margin = Pt(54)
        section.right_margin = Pt(54)

    lines = text.splitlines()
    for index, line in enumerate(lines):
        if index == 0 and line.startswith("# "):
            title = document.add_paragraph()
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = title.add_run(line[2:].strip())
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
            continue
        add_paragraph_from_markdown(document, line)

    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = generate()
    print(f"Generated {path}")
